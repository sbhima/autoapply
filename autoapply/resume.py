"""Resume text extraction + parsing into a profile dict."""
from __future__ import annotations

import json
import re
from pathlib import Path

from .llm import LLM, clip_text


def extract_text(path: Path) -> str:
    suf = path.suffix.lower()
    if suf == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    if suf == ".docx":
        import docx

        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    return path.read_text(errors="ignore")


MONTHS = "jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec"
DEGREE_RE = re.compile(
    r"\b(B\.?S\.?|B\.?A\.?|B\.?Sc|M\.?S\.?|M\.?Eng|M\.?A\.?|Ph\.?D|Bachelor(?:'s)?|Master(?:'s)?|Doctor(?:ate)?|Associate(?:'s)?)\b[^\n,|]*",
    re.I,
)
SCHOOL_RE = re.compile(r"[^\n,|•]*\b(University|College|Institute of Technology|Institute|School of)\b[^\n,|•]*", re.I)


def heuristic_parse(text: str) -> dict:
    """Best-effort regex parse — used when no LLM is configured."""
    out: dict = {"personal": {}, "links": {}, "education": [], "experience": [], "skills": []}
    lines = [l.strip() for l in text.splitlines() if l.strip()]

    m = re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)
    if m:
        out["personal"]["email"] = m.group(0)
    m = re.search(r"(\+?1[\s.-]?)?\(?\d{3}\)?[\s.-]?\d{3}[\s.-]?\d{4}", text)
    if m:
        out["personal"]["phone"] = m.group(0).strip()
    m = re.search(r"(https?://)?(www\.)?linkedin\.com/in/[\w%-]+/?", text, re.I)
    if m:
        out["links"]["linkedin"] = ("https://" + m.group(0)) if not m.group(0).startswith("http") else m.group(0)
    m = re.search(r"(https?://)?(www\.)?github\.com/[\w-]+/?", text, re.I)
    if m:
        out["links"]["github"] = ("https://" + m.group(0)) if not m.group(0).startswith("http") else m.group(0)
    for m in re.finditer(r"(https?://)?[\w-]+\.(dev|io|me|com|net|org|app|site)(/[\w./-]*)?", text, re.I):
        s = m.group(0)
        if any(x in s.lower() for x in ("linkedin", "github", "@", "gmail", "umd.edu")):
            continue
        if "." in s and not re.search(r"\.(com|net|org)$", s.split("/")[0]) or s.lower().startswith("http"):
            out["links"]["portfolio"] = s if s.startswith("http") else "https://" + s
            break

    # Name: first line that looks like 2-4 capitalised words without digits/@
    for l in lines[:5]:
        if "@" in l or re.search(r"\d", l):
            continue
        words = re.split(r"\s+", l)
        if 2 <= len(words) <= 4 and all(w[:1].isupper() for w in words if w):
            out["personal"]["first_name"] = words[0].title() if words[0].isupper() else words[0]
            out["personal"]["last_name"] = words[-1].title() if words[-1].isupper() else words[-1]
            break

    edu: dict = {}
    m = SCHOOL_RE.search(text)
    if m:
        edu["school"] = m.group(0).strip(" -–—•")
    m = DEGREE_RE.search(text)
    if m:
        deg = m.group(0).strip()
        edu["degree"] = deg
        mm = re.search(r"\b(?:in|of)\s+([A-Z][\w &]+)", deg)
        if mm:
            edu["major"] = mm.group(1).strip()
    m = re.search(r"GPA[:\s]*([0-4]\.\d{1,2})", text, re.I)
    if m:
        edu["gpa"] = m.group(1)
    m = re.search(rf"(expected|graduat\w*|anticipated)[^\n]*?(({MONTHS})\w*\.?\s+)?(20\d\d)", text, re.I)
    if m:
        edu["end"] = m.group(4) + (f"-{_month_num(m.group(3))}" if m.group(3) else "")
    if edu:
        out["education"].append(edu)

    m = re.search(r"(?:technical )?skills?\s*[:\n](.{0,600})", text, re.I | re.S)
    if m:
        chunk = m.group(1).split("\n\n")[0]
        skills = [s.strip(" •-\t") for s in re.split(r"[,•|\n;]", chunk) if 1 < len(s.strip()) < 30]
        out["skills"] = [s for s in skills if ":" not in s][:40]
    return out


def _month_num(m: str) -> str:
    keys = ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"]
    return f"{keys.index(m.lower()[:3]) + 1:02d}"


LLM_SCHEMA = """{
 "personal": {"first_name":"","last_name":"","middle_name":"","email":"","phone":"","location":"City, ST",
              "address":{"city":"","state":"","zip":"","country":""}},
 "links": {"linkedin":"","github":"","portfolio":"","twitter":""},
 "education": [{"school":"","degree":"e.g. Bachelor of Science","major":"","minor":"","gpa":"","start":"YYYY-MM","end":"YYYY-MM"}],
 "experience": [{"company":"","title":"","location":"","start":"YYYY-MM","end":"YYYY-MM or present","description":"1-2 sentences"}],
 "skills": ["..."],
 "preferences": {"years_experience": "number as string (full-time professional only, internships count as 0 unless stated)"}
}"""


def llm_parse(text: str, llm: LLM) -> dict:
    sys = "You extract structured data from resumes. Never invent data; use empty strings for unknown fields."
    prompt = f"Extract this resume into JSON matching exactly this schema:\n{LLM_SCHEMA}\n\nRESUME:\n{clip_text(text, 9000)}"
    return llm.complete_json(sys, prompt, max_tokens=1500)


def parse_resume(path: Path, llm: LLM | None) -> tuple[str, dict]:
    text = extract_text(path)
    data = heuristic_parse(text)
    if llm:
        try:
            ld = llm_parse(text, llm)
            # LLM wins where it produced something, heuristics fill gaps.
            data = _merge_nonempty(data, ld)
        except Exception as e:  # noqa: BLE001
            print(f"[resume] LLM parse failed ({e}); using heuristic parse")
    return text, data


def _merge_nonempty(a: dict, b: dict) -> dict:
    out = dict(a)
    for k, v in (b or {}).items():
        if isinstance(v, dict):
            out[k] = _merge_nonempty(out.get(k, {}) if isinstance(out.get(k), dict) else {}, v)
        elif v not in ("", None, [], {}):
            out[k] = v
    return out


if __name__ == "__main__":  # pragma: no cover
    import sys

    t, d = parse_resume(Path(sys.argv[1]), LLM.from_env())
    print(json.dumps(d, indent=2))
